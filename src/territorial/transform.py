"""Tratamento e padronização: transforma as fontes brutas em tabelas limpas.

Saída:
- ``municipios``: um polígono por município (SIRGAS 2000), com código IBGE de 7
  dígitos, nome, UF e área em hectares;
- ``cobertura``: formato "longo" (município × ano × classe → hectares), pronto para
  o Power BI, SQL ou qualquer ferramenta.
"""
from __future__ import annotations

import json
import logging
import re
from pathlib import Path

import geopandas as gpd
import pandas as pd

from . import config

log = logging.getLogger(__name__)


class ErroDeFormato(ValueError):
    """A fonte mudou de formato de um jeito que o pipeline não sabe tratar."""


# ------------------------------------------------------------------- IBGE
def montar_municipios(malha_geojson: Path, localidades_json: Path, uf: str) -> gpd.GeoDataFrame:
    gdf = gpd.read_file(malha_geojson)
    if "codarea" not in gdf.columns:
        raise ErroDeFormato(
            f"Malha do IBGE sem a coluna 'codarea' (colunas: {list(gdf.columns)})"
        )
    gdf = gdf.rename(columns={"codarea": "cd_mun"})[["cd_mun", "geometry"]]
    gdf["cd_mun"] = gdf["cd_mun"].astype(str).str.zfill(7)

    # A API entrega a malha já em SIRGAS 2000; declaramos o CRS explicitamente
    # quando o arquivo vem sem essa informação.
    if gdf.crs is None:
        gdf = gdf.set_crs(config.CRS_GEOGRAFICO)
    else:
        gdf = gdf.to_crs(config.CRS_GEOGRAFICO)

    with open(localidades_json, encoding="utf-8") as f:
        locs = json.load(f)
    nomes = pd.DataFrame(
        {"cd_mun": [str(x["id"]).zfill(7) for x in locs], "nm_mun": [x["nome"] for x in locs]}
    )
    gdf = gdf.merge(nomes, on="cd_mun", how="left")
    gdf["uf"] = uf
    gdf["area_ha_ibge"] = gdf.to_crs(config.CRS_AREA).geometry.area / 10_000.0
    return gdf[["cd_mun", "nm_mun", "uf", "area_ha_ibge", "geometry"]].sort_values("cd_mun").reset_index(drop=True)


# --------------------------------------------------------------- MapBiomas
def _primeira_coluna(df: pd.DataFrame, candidatas: tuple[str, ...], rotulo: str) -> str:
    minusculas = {str(c).strip().lower(): c for c in df.columns}
    for nome in candidatas:
        if nome in minusculas:
            return minusculas[nome]
    raise ErroDeFormato(
        f"Não encontrei a coluna de {rotulo}. Aceitas: {candidatas}. "
        f"Colunas do arquivo: {list(df.columns)[:15]}..."
    )


def _colunas_de_ano(df: pd.DataFrame) -> list:
    return [c for c in df.columns if re.fullmatch(r"\d{4}", str(c).strip())]


def escolher_aba(abas: list[str], preferida: str | None = None) -> str:
    """Aba da cobertura: a configurada, ou a primeira que comece com COVERAGE."""
    preferida = preferida or config.MAPBIOMAS_ABA
    if preferida in abas:
        return preferida
    candidatas = [a for a in abas if str(a).upper().startswith("COVERAGE")]
    if candidatas:
        log.warning("Aba %r não existe; usando %r", preferida, candidatas[0])
        return candidatas[0]
    raise ErroDeFormato(f"Aba de cobertura não encontrada. Esperada {preferida!r}; abas do arquivo: {abas}")


def _nome_de_coluna(v) -> str | int:
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return str(v).strip() if v is not None else ""


def _abrir_xlsx(arquivo: Path) -> Path:
    """Devolve o caminho de um .xlsx de verdade (o Drive pode entregar a planilha dentro de um .zip)."""
    import zipfile

    if not zipfile.is_zipfile(arquivo):
        return arquivo
    with zipfile.ZipFile(arquivo) as z:
        nomes = z.namelist()
        if "[Content_Types].xml" in nomes:
            return arquivo  # já é um .xlsx
        internos = [n for n in nomes if n.lower().endswith((".xlsx", ".xlsm")) and not n.endswith("/")]
        if not internos:
            raise ErroDeFormato(
                "O arquivo do MapBiomas é um zip sem nenhuma planilha .xlsx dentro "
                f"(conteúdo: {', '.join(nomes[:5])})."
            )
        interno = max(internos, key=lambda n: z.getinfo(n).file_size)
        destino_dir = arquivo.parent / "descompactado"
        destino_dir.mkdir(parents=True, exist_ok=True)
        destino = destino_dir / Path(interno).name
        if not destino.exists() or destino.stat().st_size != z.getinfo(interno).file_size:
            log.info("MapBiomas: descompactando %s", interno)
            with z.open(interno) as src, open(destino, "wb") as dst:
                while bloco := src.read(1024 * 1024):
                    dst.write(bloco)
        return destino


def ler_mapbiomas(arquivo_xlsx: Path, uf: str, aba: str | None = None) -> pd.DataFrame:
    """Lê a tabela larga do MapBiomas e devolve o formato longo para a UF pedida.

    A planilha do Brasil inteiro é grande. Por isso a leitura é em *streaming* (linha a
    linha, modo somente-leitura do openpyxl) e só as linhas da UF ficam na memória —
    o consumo não cresce com o tamanho do arquivo.
    """
    from openpyxl import load_workbook

    prefixo = config.ESTADOS[uf].codigo_ibge
    wb = load_workbook(_abrir_xlsx(Path(arquivo_xlsx)), read_only=True, data_only=True)
    try:
        ws = wb[escolher_aba(wb.sheetnames, aba)]
        linhas = ws.iter_rows(values_only=True)
        try:
            cabecalho = [_nome_de_coluna(c) for c in next(linhas)]
        except StopIteration:
            raise ErroDeFormato("A aba de cobertura do MapBiomas está vazia.") from None
        # Acha a coluna do município com as mesmas regras da conversão (erro claro se faltar).
        col_mun = _primeira_coluna(pd.DataFrame(columns=cabecalho), config.MAPBIOMAS_COLUNAS_ID_MUNICIPIO,
                                   "código do município")
        i_mun = cabecalho.index(col_mun)
        mantidas, total = [], 0
        for linha in linhas:
            total += 1
            if linha is None or len(linha) <= i_mun or linha[i_mun] is None:
                continue
            codigo = str(linha[i_mun]).strip()
            if codigo.endswith(".0"):
                codigo = codigo[:-2]
            if codigo.zfill(7).startswith(prefixo):
                mantidas.append(linha[: len(cabecalho)])
        log.info("MapBiomas: %d linhas lidas, %d da UF %s", total, len(mantidas), uf)
    finally:
        wb.close()
    bruto = pd.DataFrame(mantidas, columns=cabecalho)
    return mapbiomas_para_longo(bruto, uf)


def mapbiomas_para_longo(bruto: pd.DataFrame, uf: str) -> pd.DataFrame:
    col_mun = _primeira_coluna(bruto, config.MAPBIOMAS_COLUNAS_ID_MUNICIPIO, "código do município")
    col_classe = _primeira_coluna(bruto, config.MAPBIOMAS_COLUNAS_CLASSE_ID, "id da classe")
    anos = _colunas_de_ano(bruto)
    if not anos:
        raise ErroDeFormato("Nenhuma coluna de ano (AAAA) na tabela do MapBiomas.")

    # nome da classe é opcional (só enriquece a tabela)
    minusculas = {str(c).strip().lower(): c for c in bruto.columns}
    col_nome = next((minusculas[n] for n in config.MAPBIOMAS_COLUNAS_CLASSE_NOME if n in minusculas), None)

    prefixo = config.ESTADOS[uf].codigo_ibge
    df = bruto.copy()
    df["cd_mun"] = df[col_mun].astype(str).str.replace(r"\.0$", "", regex=True).str.zfill(7)
    df = df[df["cd_mun"].str.startswith(prefixo)]

    chaves = ["cd_mun", col_classe] + ([col_nome] if col_nome else [])
    longo = df.melt(id_vars=chaves, value_vars=anos, var_name="ano", value_name="area_ha")
    longo = longo.rename(columns={col_classe: "classe_id", **({col_nome: "classe_nome"} if col_nome else {})})
    if "classe_nome" not in longo.columns:
        longo["classe_nome"] = None
    longo["ano"] = longo["ano"].astype(str).str.strip().astype(int)
    longo["classe_id"] = pd.to_numeric(longo["classe_id"], errors="coerce")
    longo["area_ha"] = pd.to_numeric(longo["area_ha"], errors="coerce")

    # Um município pode aparecer em mais de um bioma: somamos por município/ano/classe.
    longo = (
        longo.groupby(["cd_mun", "ano", "classe_id"], as_index=False, dropna=False)
        .agg(area_ha=("area_ha", "sum"), classe_nome=("classe_nome", "first"))
        .sort_values(["cd_mun", "ano", "classe_id"])
        .reset_index(drop=True)
    )
    return longo


# ---------------------------------------------------------------- derivados
def resumo_por_municipio(municipios: gpd.GeoDataFrame, cobertura: pd.DataFrame, classe_id: int) -> pd.DataFrame:
    """% de uma classe (ex.: floresta) no último ano disponível, por município."""
    ultimo = int(cobertura["ano"].max())
    c = cobertura[cobertura["ano"] == ultimo]
    total = c.groupby("cd_mun")["area_ha"].sum().rename("area_total_ha")
    alvo = c[c["classe_id"] == classe_id].groupby("cd_mun")["area_ha"].sum().rename("area_classe_ha")
    r = pd.concat([total, alvo], axis=1).fillna(0.0).reset_index()
    r["pct_classe"] = (r["area_classe_ha"] / r["area_total_ha"].where(r["area_total_ha"] > 0)) * 100
    r["ano"] = ultimo
    return municipios[["cd_mun", "nm_mun"]].merge(r, on="cd_mun", how="left")
